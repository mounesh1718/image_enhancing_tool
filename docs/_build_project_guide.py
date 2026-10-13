"""Build the project guide from the current Image Quality Enhancer implementation."""

from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
ART = ROOT / ".guide_assets"
ART.mkdir(parents=True, exist_ok=True)
OUT = ROOT / "ImageQualityEnhancer_Project_Guide.docx"

NAVY = "17324D"
BLUE = "2877A9"
TEAL = "157D82"
MINT = "EAF5F3"
PALE = "EEF4F8"
GREY = "5B6670"
LINE = "D5DFE6"
INK = "1F2D38"
AMBER = "A76818"


def font(size=30, bold=False):
    base = Path("C:/Windows/Fonts")
    name = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(base / name), size)


def wrapped(draw, value, face, width):
    words = value.split()
    rows = []
    row = ""
    for word in words:
        trial = (row + " " + word).strip()
        if row and draw.textbbox((0, 0), trial, font=face)[2] > width:
            rows.append(row)
            row = word
        else:
            row = trial
    if row:
        rows.append(row)
    return rows


def center_rows(draw, rows, y, face, fill, cx, step):
    for row in rows:
        box = draw.textbbox((0, 0), row, font=face)
        draw.text((cx - (box[2] - box[0]) / 2, y), row, font=face, fill=fill)
        y += step


def flow_card(draw, box, title, detail, tech, bg="white", accent=BLUE):
    x0, y0, x1, y1 = box
    if bg not in ("white", "black") and not bg.startswith("#"):
        bg = "#" + bg
    draw.rounded_rectangle(box, radius=22, fill=bg, outline="#D5DFE6", width=3)
    draw.rounded_rectangle((x0, y0, x0 + 16, y1), radius=8, fill="#" + accent)
    f1, f2, f3 = font(36, True), font(28), font(24, True)
    draw.text((x0 + 43, y0 + 24), title, font=f1, fill="#" + NAVY)
    rows = wrapped(draw, detail, f2, x1 - x0 - 85)
    for j, row in enumerate(rows[:2]):
        draw.text((x0 + 43, y0 + 78 + j * 34), row, font=f2, fill="#" + INK)
    draw.text((x0 + 43, y1 - 38), "TECH  " + tech, font=f3, fill="#" + accent)


def arrow(draw, cx, y1, y2):
    draw.line((cx, y1, cx, y2 - 12), fill="#7B96A8", width=6)
    draw.polygon([(cx - 12, y2 - 17), (cx + 12, y2 - 17), (cx, y2)], fill="#7B96A8")


def make_main_flow():
    im = Image.new("RGB", (1900, 1550), "white")
    d = ImageDraw.Draw(im)
    cards = [
        ("1  Upload and copy", "PDF, DOCX or a standalone image enters a local job workspace.", "Avalonia UI · C# job orchestrator"),
        ("2  Find and extract pictures", "Record each placement; extract each unique raster picture once without changing its pixels.", "Open XML SDK + ZIP · PyMuPDF · Python image I/O"),
        ("3  Sort by image type", "Classify chemical, chart/flow, table or other; inspect mixed panels; ask the user when unsure.", "DocumentFigureClassifier v2.5 · ONNX Runtime · MolScribe check"),
        ("4  Measure and improve", "Read pixels, measure blur/noise/contrast, then apply conservative repair only where indicated.", "OpenCV · NumPy · project rules"),
        ("5  Optional specialist branch", "On request, try suitable AI repair; chemical redraw has a separate structure-verification path.", "SCUNet · FBCNN · DocRes · Real-ESRGAN · MolScribe · RDKit"),
        ("6  Compare and approve", "Check text, image detail and chemistry as relevant. Show original and candidates side by side.", "RapidOCR · image metrics · human reviewer"),
        ("7  Export new copy", "Replace approved pictures only; validate untouched content and preserve the original file.", "DOCX ZIP rewrite / Open XML · PyMuPDF · HTML report"),
    ]
    top, height, gap = 45, 175, 39
    for i, (title, detail, tech) in enumerate(cards):
        y = top + i * (height + gap)
        flow_card(d, (90, y, 1810, y + height), title, detail, tech,
                  bg="#F3F8FB" if i % 2 == 0 else "white",
                  accent=TEAL if i in (2, 5) else BLUE)
        if i < len(cards) - 1:
            arrow(d, 950, y + height + 3, y + height + gap - 3)
    p = ART / "workflow.png"
    im.save(p, dpi=(240, 240))
    return p


def make_enhancement_flow():
    im = Image.new("RGB", (1900, 1000), "white")
    d = ImageDraw.Draw(im)
    flow_card(d, (80, 35, 1820, 185), "Extracted raster picture", "Original pixels stay available throughout the job.", "DOCX/PDF extractor or standalone image", PALE)
    arrow(d, 950, 187, 235)
    flow_card(d, (80, 240, 1820, 390), "Measure", "Blur, grain/noise, sharpness, tones, text and edges become numeric readings.", "OpenCV + NumPy + Python decision thresholds", "white")
    arrow(d, 950, 392, 440)
    flow_card(d, (80, 445, 1820, 595), "Normal enhancement", "Conditional denoise, background/contrast repair and careful sharpening; clear images stay unchanged.", "OpenCV filters; same output dimensions", PALE)
    arrow(d, 950, 597, 645)
    flow_card(d, (80, 650, 1820, 800), "Optional AI after review-screen click", "Choose a model that fits the measured defect; build a separate candidate from the original.", "SCUNet / FBCNN / DocRes / Real-ESRGAN", "white", TEAL)
    arrow(d, 950, 802, 850)
    flow_card(d, (80, 855, 1820, 990), "Safety checks and user choice", "Compare original, OpenCV and AI candidates; reject harmful changes before approval.", "RapidOCR + detail metrics + UI review", PALE, TEAL)
    p = ART / "enhancement.png"
    im.save(p, dpi=(240, 240))
    return p


def make_chem_flow():
    im = Image.new("RGB", (1900, 1250), "white")
    d = ImageDraw.Draw(im)
    cards = [
        ("Chemical candidate only", "Classifier puts the image in Chemical diagrams; MolScribe is the second reading, not the initial extractor.", "DocumentFigureClassifier · MolScribe"),
        ("Read the drawing", "Predict atom labels, bond connections, coordinates and wedge/dash stereo marks.", "Pretrained MolScribe model"),
        ("Build a molecule", "Expand known short labels; check valence/connectivity and calculate SMILES, formula and InChIKey.", "RDKit chemical toolkit"),
        ("Draw and verify", "Redraw at the original size/layout; check ink-to-bond alignment, OCR labels, wedges and a second reading.", "RDKit drawing · OpenCV · RapidOCR · MolScribe"),
        ("Read back, then review", "The redrawn picture must read as the same molecule, including stereo. A chemist still approves.", "MolScribe · RDKit InChIKey · human chemist"),
    ]
    top, height, gap = 35, 190, 42
    for i, (title, detail, tech) in enumerate(cards):
        y = top + i * (height + gap)
        flow_card(d, (80, y, 1820, y + height), title, detail, tech,
                  bg="#F1F8F7" if i % 2 == 0 else "white", accent=TEAL)
        if i < len(cards) - 1:
            arrow(d, 950, y + height + 2, y + height + gap - 2)
    p = ART / "chemistry.png"
    im.save(p, dpi=(240, 240))
    return p


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def margins(cell, top=85, start=100, bottom=85, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for key, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement("w:" + key)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        m.append(node)
    tcPr.append(m)


def set_run(run, size=9.5, bold=False, color=INK):
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def paragraph(doc, text="", style=None, size=9.5, bold=False, color=INK, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    set_run(p.add_run(text), size, bold, color)
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    h.paragraph_format.space_after = Pt(5)
    for r in h.runs:
        set_run(r, 16 if level == 1 else 11, True, NAVY)
    return h


def bullet(doc, text):
    return paragraph(doc, text, "List Bullet", 9.2, after=3)


def table(doc, headers, rows, widths=None, font_size=8.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.style = "Table Grid"
    if widths:
        for c, w in zip(t.rows[0].cells, widths):
            c.width = Inches(w)
    for c, value in zip(t.rows[0].cells, headers):
        shade(c, NAVY)
        margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c.text = ""
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(value), font_size, True, "FFFFFF")
    for idx, row in enumerate(rows):
        cells = t.add_row().cells
        for c, value in zip(cells, row):
            if idx % 2 == 0:
                shade(c, "F4F7F9")
            margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run(p.add_run(value), font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def pic(doc, path, width=7.05):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.add_run().add_picture(str(path), width=Inches(width))


def newpage(doc):
    doc.add_page_break()


def section_kicker(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run(text.upper()), 8.5, True, TEAL)


def build():
    wf, ef, cf = make_main_flow(), make_enhancement_flow(), make_chem_flow()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = Inches(.72)
    sec.bottom_margin = Inches(.66)
    sec.left_margin = sec.right_margin = Inches(.70)
    sec.header_distance = Inches(.28)
    sec.footer_distance = Inches(.30)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for n in ("Title", "Heading 1", "Heading 2"):
        styles[n].font.name = "Aptos Display" if n == "Title" else "Aptos"
        styles[n].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Title"].font.size = Pt(29)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(11)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(footer.add_run("IMAGE QUALITY ENHANCER  •  PROJECT GUIDE  |  "), 7.5, False, GREY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    # 1 — opening
    section_kicker(doc, "Project overview  |  October 2026")
    t = doc.add_paragraph(style="Title")
    t.paragraph_format.space_after = Pt(8)
    t.add_run("Image Quality Enhancer")
    paragraph(doc, "How document pictures are extracted, classified, improved, checked and safely returned", size=13, color=BLUE, after=17)
    heading(doc, "What this tool does")
    paragraph(doc, "Image Quality Enhancer is an offline Windows desktop app for improving raster pictures inside PDF and DOCX files, or standalone image files. It sorts pictures into useful groups, applies conservative image repair, offers optional specialist AI, and lets a person compare each result before saving a new document. The original file is never overwritten.", size=10.4, after=11)
    heading(doc, "The core idea")
    table(doc, ["Input", "Decision and work", "Output"], [
        ("PDF, DOCX or image", "Find each picture, measure its quality and identify its type", "A reviewable OpenCV result or unchanged original"),
        ("Optional user action", "Try one suitable AI repair, or request chemical redraw", "A separate candidate with explicit safety checks"),
        ("Human approval", "Choose the version to keep for each picture", "A validated new file and an HTML report"),
    ], font_size=9.2)
    heading(doc, "Four picture groups")
    paragraph(doc, "Chemical diagrams  •  Flowcharts and graphs  •  Tables  •  Other images, including photographs, logos, screenshots and maps. Mixed figure sheets can appear in more than one group. Classification is a suggestion, not an infallible verdict.", size=10, after=12)
    heading(doc, "What this guide covers")
    paragraph(doc, "The following pages show the implemented processing path, the distinct chemical and optional AI routes, the review safeguards, and a technology register with licence and pretrained-model status. Names and behaviour were checked against this project's README and source code; model and library licences were checked against upstream project pages. Licence summaries are informational, not legal advice.", size=9.4)

    # 2 — end to end
    newpage(doc)
    section_kicker(doc, "Process map")
    heading(doc, "The complete workflow")
    paragraph(doc, "Read the diagram from top to bottom. Each box names the process and the technology that performs it.", after=8)
    pic(doc, wf, 7.05)
    paragraph(doc, "Normal image repair happens first. The specialist AI buttons are available on the review screen afterward; a model runs only when the user asks. Chemical redraw is a separate path selected at job start and only applies to chemical candidates.", size=9.2, color=GREY)

    # 3 — extraction and enhancement
    newpage(doc)
    section_kicker(doc, "Input and image quality")
    heading(doc, "From document to improved picture")
    table(doc, ["Source", "How pictures are taken out", "What is preserved"], [
        ("DOCX", "Open XML SDK locates picture parts; ZIP handling copies their bytes.", "Each picture's original pixels, placement and layout."),
        ("PDF", "PyMuPDF finds and extracts image streams; supported images are decoded for processing.", "JPEG streams where possible; untouched PDF objects stay unchanged."),
        ("Image file", "The original file is read directly.", "The uploaded original remains available."),
    ], font_size=8.6)
    paragraph(doc, "Vector graphics such as SVG, EMF, WMF and PDF vector drawings are kept unchanged. Some unusual images are deliberately skipped with a reason, including inline PDF images, CMYK/spot-colour sources, high-bit-depth images, linked images, OLE previews and images over limits.", size=9.2, after=8)
    pic(doc, ef, 7.05)
    heading(doc, "What the measurements mean", 2)
    paragraph(doc, "OpenCV decodes pixels into NumPy arrays. Project code calculates numeric clues such as sharpness, blur, grain/noise, contrast, tonal range and edge preservation. Its thresholds choose gentle OpenCV operations; NumPy does not decide or improve the picture by itself. A strong result is one that improves readability without losing lines, characters, colours or chemical detail.", size=9.2)

    # 4 — options
    newpage(doc)
    section_kicker(doc, "Routing and optional models")
    heading(doc, "Which specialist model is tried")
    paragraph(doc, "After ordinary OpenCV processing, the review screen suggests a model only when the measured problem fits. The user can try it for one picture or all eligible pictures. A separate AI candidate is made from the original, then compared with the OpenCV result.", size=9.5, after=8)
    table(doc, ["Measured problem", "Model", "What it may improve", "Where used"], [
        ("Scan or camera grain", "SCUNet", "Reduce random grain while retaining detail", "Eligible raster images"),
        ("Strong JPEG damage", "FBCNN", "Reduce block edges and ringing caused by compression", "Images saved as low-quality JPEG"),
        ("Blurred readable text", "DocRes", "Make text strokes clearer", "Text-bearing images only"),
        ("Soft or noisy photo", "Real-ESRGAN", "Restore photo texture and apparent clarity", "Photos in Other images only"),
    ], font_size=8.5)
    heading(doc, "How an AI candidate is accepted", 2)
    bullet(doc, "RapidOCR compares readable text in the original and candidate. Changed digits or too many changed characters stop the candidate; DocRes has stricter checks for newly appearing text.")
    bullet(doc, "Image-detail checks compare lines and edges with the OpenCV candidate. The AI version cannot add or lose too much detail.")
    bullet(doc, "For chemical pictures, MolScribe and RDKit must read the same molecular identity before and after enhancement.")
    bullet(doc, "A passing AI version still needs human approval. The reviewer can switch back to the OpenCV result or keep the original.")
    heading(doc, "Classification happens before specialist processing", 2)
    paragraph(doc, "DocumentFigureClassifier v2.5 is a pretrained EfficientNet-B0 model run by ONNX Runtime. It scores 26 figure types and the app groups them into four UI categories. For a possible chemical picture, MolScribe gives a second structure reading. Low-confidence categories remain visible for user correction; tables and photos do not receive chemical error messages.", size=9.3)

    # 5 — chemistry
    newpage(doc)
    section_kicker(doc, "Chemical diagrams")
    heading(doc, "Chemical redraw is a separate safety path")
    pic(doc, cf, 7.05)
    heading(doc, "What the chemistry checks can and cannot prove", 2)
    paragraph(doc, "MolScribe is the pretrained AI reader: it estimates atoms, bonds and stereochemical wedge/dash marks from the picture. RDKit is not an AI model and does not consult a catalogue of known chemical names. It tests whether the interpreted molecular graph is chemically consistent, generates identifiers such as SMILES and InChIKey, and helps draw the molecule. A newly invented but valid structure can be handled; a valid RDKit molecule does not prove that the AI read the original picture perfectly.", size=9.1)
    paragraph(doc, "The redraw tries to keep the source layout, condensed labels and original pixel size. Automatic checks compare the drawing against the reading, then read the redraw back. The app offers Improve only, Redraw when verified and Redraw all. Unverified redraws are labelled clearly; every actual redraw requires a chemist's approval and is excluded from bulk safe approval.", size=9.1)

    # 6 — safety and outputs
    newpage(doc)
    section_kicker(doc, "Decision and delivery")
    heading(doc, "What the reviewer sees")
    table(doc, ["Image type", "Review focus", "Controls"], [
        ("Normal images, tables and charts", "Check text, lines, colours and legibility; no irrelevant chemistry warning.", "Original ↔ OpenCV, slider or side by side; Approve or Keep original."),
        ("Optional AI result", "Compare original, OpenCV and AI candidate; read any rejection reason.", "Choose version; approve each AI result individually."),
        ("Chemical diagram", "Check every atom, bond, label and wedge/dash stereo mark; inspect numbered trouble areas.", "Manual chemist approval for any redraw."),
    ], font_size=8.8)
    heading(doc, "Safety rules built into the workflow")
    bullet(doc, "Original files are copied into a local job workspace and are never overwritten.")
    bullet(doc, "Enhanced images keep the original pixel dimensions; already-clear images may remain unchanged.")
    bullet(doc, "Only approved candidates are inserted into a new DOCX or PDF. Untouched document parts are validated; output is published only after validation.")
    bullet(doc, "A bad or unsupported image keeps its original. If a worker fails, the job can retry that image without losing the rest of the work.")
    bullet(doc, "AI is optional and installed once; after setup the app runs offline. The reviewer remains responsible for scientific correctness.")
    heading(doc, "Delivered result")
    paragraph(doc, "The user saves a new document or image file and can save an HTML report with before/after previews, locations, measurements, selected version, checks and decisions. A changed signed PDF requires explicit confirmation because the new copy's signature will no longer be valid.", size=9.6)
    heading(doc, "Important limits")
    paragraph(doc, "Image similarity, OCR and matching InChIKeys are useful safeguards, not a promise of 100% accuracy. OCR can misread an already-blurred original, image models can invent texture, and a chemical model can misread a bond. For regulated or research documents, a qualified person should compare the output with the source at full size.", size=9.5)

    # 7 — register
    newpage(doc)
    section_kicker(doc, "Technology register")
    heading(doc, "Core software and their purpose")
    paragraph(doc, "Pretrained means a ready-to-use learned model exists. Ordinary libraries and frameworks do not have a training state.", size=9, after=7)
    table(doc, ["Technology", "Purpose in this project", "Licence", "Pretrained"], [
        (".NET 10, C#; Avalonia", "Desktop UI and background job control", "MIT; MIT", "No"),
        ("CommunityToolkit.Mvvm", "UI state and commands", "MIT", "No"),
        ("Open XML SDK; ZIP handling", "Locate DOCX images; preserve/rewrite package parts", "MIT; Python stdlib", "No"),
        ("PyMuPDF", "Scan/extract PDF images and validate PDF output", "AGPL-3.0 or commercial", "No"),
        ("OpenCV", "Decode, measure and conservatively filter pixels", "Apache-2.0", "No"),
        ("NumPy", "Pixel arrays and numeric measurements", "BSD-3-Clause*", "No"),
        ("ONNX Runtime / DirectML", "Run classifier, OCR and image models on CPU/GPU", "MIT", "No"),
        ("RapidOCR + PP-OCRv5", "Read text before/after enhancement; read atom labels", "Apache-2.0†", "Yes, OCR weights"),
        ("RDKit", "Validate interpreted molecular graph, identifiers and redraw", "BSD-3-Clause", "No"),
    ], font_size=8.1)
    paragraph(doc, "* NumPy binary distributions can contain separately licensed components. † OCR package and upstream model notices should be shipped with a distribution. PyMuPDF licensing is a significant distribution decision: review AGPL obligations or obtain a commercial licence before proprietary release.", size=8.7, color=GREY)
    heading(doc, "Pretrained models used by the project")
    table(doc, ["Model", "Purpose", "Licence", "Pretrained"], [
        ("DocumentFigureClassifier v2.5", "Sort 26 figure types into four review groups", "MIT", "Yes; supplied checkpoint"),
        ("MolScribe", "Read atoms, bonds and stereo from chemical pictures", "MIT‡", "Yes; supplied checkpoint"),
        ("SCUNet", "Remove image grain", "Apache-2.0", "Yes; author weights"),
        ("FBCNN", "Repair JPEG compression damage", "Apache-2.0", "Yes; author weights"),
        ("DocRes", "Restore blurred document text", "MIT‡", "Yes; author weights"),
        ("Real-ESRGAN general x4v3", "Restore soft or noisy photos", "BSD-3-Clause", "Yes; author weights"),
    ], font_size=8.1)
    paragraph(doc, "‡ The project uses a pinned MolScribe model repository and a hash-pinned DocRes weight copy; verify the exact weight-file terms and attribution again when redistributing model files. None of these models needs training by this project's user. Installation downloads optional checkpoints once; runtime processing remains local.", size=8.7, color=GREY)

    # 8 — references and ownership
    newpage(doc)
    section_kicker(doc, "Sources and implementation notes")
    heading(doc, "Where the details were checked")
    paragraph(doc, "Project source: README.md; ImageQualityEnhancer.csproj; Services/JobOrchestrator.cs; Services/Docx; worker/figure_classifier.py; worker/analyse.py; worker/enhance.py; worker/ai_assist.py; worker/textcheck.py; worker/redraw; worker/pdf_ops.py; worker/setup_ai.py. This document describes the current implementation, not a proposal for future functionality.", size=9.2, after=12)
    heading(doc, "Upstream model and licence references")
    refs = [
        "DocumentFigureClassifier v2.5 — huggingface.co/docling-project/DocumentFigureClassifier-v2.5",
        "MolScribe — github.com/thomas0809/MolScribe and huggingface.co/yujieq/MolScribe",
        "RDKit — github.com/rdkit/rdkit/blob/master/license.txt",
        "SCUNet — github.com/cszn/SCUNet  |  FBCNN — github.com/jiaxi-jiang/FBCNN",
        "DocRes — github.com/ZZZHANG-jx/DocRes  |  Real-ESRGAN — github.com/xinntao/Real-ESRGAN",
        "RapidOCR — github.com/RapidAI/RapidOCR  |  ONNX Runtime — github.com/microsoft/onnxruntime",
        "PyMuPDF licensing — pymupdf.readthedocs.io/en/latest/about.html",
        "Open XML SDK — github.com/dotnet/Open-XML-SDK  |  Avalonia — github.com/AvaloniaUI/Avalonia",
    ]
    for ref in refs:
        bullet(doc, ref)
    heading(doc, "Operational ownership")
    paragraph(doc, "The developer owns the measured thresholds, category routing, model setup and output validation. The app proposes changes but never certifies scientific accuracy. The user owns final visual approval, and a chemist approves every redraw. Before distributing a packaged product, confirm all library/model notices and especially the PyMuPDF licence choice.", size=9.5)
    heading(doc, "Bottom line")
    paragraph(doc, "This is a review-first document image enhancer. It uses measured pixel repair for every eligible picture, specialist pretrained models only when requested and appropriate, and extra graph/stereochemistry checks for chemical diagrams. The final new file contains only the versions a person approved.", size=10.2, bold=True, color=NAVY)

    doc.core_properties.title = "Image Quality Enhancer Project Guide"
    doc.core_properties.subject = "Workflow, technology, pretrained models and licences"
    doc.core_properties.author = "Image Quality Enhancer Project"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
